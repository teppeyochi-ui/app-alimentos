#!/usr/bin/env python3
"""Gera Teste_Classificador.xlsx a partir do JSON gravado por teste-motor.js.

Uso (na pasta classificador-camarao):
    node teste-motor.js --json ferramentas/relatorio-testes.json
    python3 ferramentas/gerar-planilha.py ferramentas/relatorio-testes.json Teste_Classificador.xlsx
"""
import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FONTE = 'Arial'
CAB = PatternFill('solid', fgColor='1F4E79')
OK = PatternFill('solid', fgColor='E2EFDA')
FALHA = PatternFill('solid', fgColor='F8CBAD')


def aba(wb, titulo, cabecalho, linhas, larguras=None):
    ws = wb.create_sheet(titulo)
    ws.append(cabecalho)
    for c in ws[1]:
        c.font = Font(name=FONTE, bold=True, color='FFFFFF')
        c.fill = CAB
        c.alignment = Alignment(vertical='center', wrap_text=True)
    for linha in linhas:
        ws.append(linha)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name=FONTE)
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    for i, _ in enumerate(cabecalho, 1):
        ws.column_dimensions[get_column_letter(i)].width = (larguras or {}).get(i, 16)
    return ws


def pintar_resultado(ws, col):
    for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
        for c in row:
            c.fill = OK if c.value in ('OK', 0) else FALHA


def main(entrada, saida):
    d = json.load(open(entrada, encoding='utf-8'))
    wb = Workbook()
    wb.remove(wb.active)

    # Resumo: contagens calculadas por fórmula a partir das demais abas.
    ws = wb.create_sheet('Resumo')
    ws['A1'] = 'Classificador de Camarão — relatório de testes do motor'
    ws['A1'].font = Font(name=FONTE, bold=True, size=14)
    ws['A2'] = 'Gerado em ' + d['data'][:10] + ' · motor ' + d['versaoMotor'] + ' · fonte: node teste-motor.js'
    itens = [
        ('Configurações testadas (esperado 72)', '=COUNTA(Configuracoes!A2:A1000)'),
        ('Pontos testados (esperado 2.352)', '=COUNTA(Pontos!A2:A10000)'),
        ('Falhas na varredura (esperado 0)', '=COUNTIF(Pontos!M2:M10000,"FALHA")'),
        ('Casos de referência e bordas', '=COUNTA(Casos!A2:A1000)'),
        ('Falhas em casos e bordas (esperado 0)', '=COUNTIF(Casos!E2:E1000,"FALHA")'),
        ('Achados nas tabelas', '=COUNTA(Achados!A2:A1000)'),
        ('Verificações totais do teste-motor.js', d['resumo']['verificacoes']),
        ('Falhas totais do teste-motor.js', d['resumo']['falhas']),
    ]
    for i, (rot, val) in enumerate(itens, start=4):
        ws.cell(row=i, column=1, value=rot).font = Font(name=FONTE)
        c = ws.cell(row=i, column=2, value=val)
        c.font = Font(name=FONTE, bold=True)
    ws['A13'] = 'Condições da varredura: fresco, congelado, glaciado 15% compensada, glaciado 15% não compensada; embalagem de 1 kg; mínimo, meio e máximo de cada classe.'
    ws['A14'] = 'Critério: a classe retornada contém o ponto; com sobreposição, vale a faixa mais estreita. "*" na aba Faixas = proporcional à coluna de 1 kg (ou à mais próxima).'
    for r in (2, 13, 14):
        ws.cell(row=r, column=1).font = Font(name=FONTE, italic=True, color='555555')
    ws.column_dimensions['A'].width = 48
    ws.column_dimensions['B'].width = 14

    c = d['configuracoes']
    w = aba(wb, 'Configuracoes', ['Espécie', 'Apresentação', 'Condição', 'Tabela', 'Unidade', 'Classes', 'Pontos', 'Falhas'],
            [[x['especie'], x['apresentacao'], x['condicao'], x['tabela'], x['unidade'], x['classes'], x['pontos'], x['falhas']] for x in c],
            {3: 24, 4: 40})
    pintar_resultado(w, 8)

    p = d['pontos']
    w = aba(wb, 'Pontos', ['Espécie', 'Apresentação', 'Condição', 'Tabela', 'Classe testada', 'Ponto', 'Peso líquido (g)', 'Peso pesado (g)',
                           'Classe retornada', 'Também cabe em', 'Faixa de peças (1 kg)', '', 'Resultado'],
            [[x['especie'], x['apresentacao'], x['condicao'], x['tabela'], x['classe_testada'], x['ponto'], round(x['peso_liquido_g'], 4),
              round(x['peso_pesado_g'], 4), x['classe_retornada'], x['sobrepostas'], x['faixa_pecas'], '', x['resultado']] for x in p],
            {3: 24, 4: 40, 11: 22, 12: 2})
    for row in w.iter_rows(min_row=2, min_col=7, max_col=8):
        for cell in row:
            cell.number_format = '0.00'
    pintar_resultado(w, 13)

    w = aba(wb, 'Casos', ['Caso', 'Peso pesado (g)', 'Esperado', 'Obtido', 'Resultado'],
            [[x['caso'], x['peso_pesado_g'], x['esperado'], x['obtido'], x['resultado']] for x in d['casos']],
            {1: 58, 3: 26, 4: 70})
    pintar_resultado(w, 5)

    f = d['faixasEmbalagem']
    emb = ['200 g', '300 g', '400 g', '500 g', '800 g', '1 kg', '2 kg', '5 kg']
    aba(wb, 'Faixas', ['Tabela', 'Classe', 'Unidade', 'Faixa de peso', 'Faixa congelado'] + emb,
        [[x['tabela'], x['classe'], x['unidade'], x['faixa_peso'], x.get('faixa_congelado', '')] + [x[e] for e in emb] for x in f],
        {1: 42, 4: 18, 5: 16})

    aba(wb, 'Achados', ['Tabela', 'Tipo', 'Detalhe'], [[a['tabela'], a['tipo'], a['detalhe']] for a in d['achados']], {1: 42, 2: 26, 3: 110})

    wb.save(saida)


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
